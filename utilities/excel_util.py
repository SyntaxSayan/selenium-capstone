import openpyxl
from utilities.custom_logger import CustomLogger

logger = CustomLogger.get_logger("ExcelUtil")

class ExcelUtil:
    """Utility class to read and write test data using openpyxl."""

    @staticmethod
    def get_row_count(file_path, sheet_name):
        workbook = openpyxl.load_workbook(file_path, data_only=True)
        sheet = workbook[sheet_name]
        return sheet.max_row

    @staticmethod
    def get_column_count(file_path, sheet_name):
        workbook = openpyxl.load_workbook(file_path, data_only=True)
        sheet = workbook[sheet_name]
        return sheet.max_column

    @staticmethod
    def get_cell_data(file_path, sheet_name, row_num, col_num):
        workbook = openpyxl.load_workbook(file_path, data_only=True)
        sheet = workbook[sheet_name]
        return sheet.cell(row=row_num, column=col_num).value

    @staticmethod
    def set_cell_data(file_path, sheet_name, row_num, col_num, data):
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook[sheet_name]
        sheet.cell(row=row_num, column=col_num).value = data
        workbook.save(file_path)

    @staticmethod
    def get_data_as_list_of_dicts(file_path, sheet_name):
        """Reads rows and returns a list of dictionaries with column headers as keys."""
        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            sheet = workbook[sheet_name]
            headers = [sheet.cell(row=1, column=col).value for col in range(1, sheet.max_column + 1)]
            
            data_list = []
            for row in range(2, sheet.max_row + 1):
                row_data = {}
                is_empty_row = True
                for col_idx, header in enumerate(headers, start=1):
                    val = sheet.cell(row=row, column=col_idx).value
                    if val is not None:
                        is_empty_row = False
                    row_data[str(header).strip()] = val
                if not is_empty_row:
                    data_list.append(row_data)

            logger.info(f"Loaded {len(data_list)} rows from Excel sheet '{sheet_name}' in {file_path}")
            return data_list
        except Exception as e:
            logger.error(f"Error reading Excel file {file_path}: {e}")
            raise

    @staticmethod
    def get_data_as_tuples(file_path, sheet_name):
        """Returns rows as a list of tuples (excluding header row)."""
        data_dicts = ExcelUtil.get_data_as_list_of_dicts(file_path, sheet_name)
        return [tuple(d.values()) for d in data_dicts]
