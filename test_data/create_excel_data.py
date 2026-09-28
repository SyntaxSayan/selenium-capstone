import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os

def generate_test_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PurchaseScenarios"

    # Define headers
    headers = [
        "TestID", "Description", "Email", "Password",
        "SearchTerm", "ProductName", "InitialQty", "UpdatedQty",
        "ExpectedUnitPrice", "ExpectedTotal"
    ]

    # Data rows
    rows = [
        [
            "TC_E2E_01", "MacBook Purchase Flow", "tester_capstone_selenium@gmail.com", "Password@123",
            "MacBook", "MacBook", 1, 2, 602.00, 1204.00
        ],
        [
            "TC_E2E_02", "iPhone Purchase Flow", "tester_capstone_selenium@gmail.com", "Password@123",
            "iPhone", "iPhone", 1, 3, 123.20, 369.60
        ],
        [
            "TC_E2E_03", "Samsung Tab Purchase Flow", "tester_capstone_selenium@gmail.com", "Password@123",
            "Samsung Galaxy", "Samsung Galaxy Tab 10.1", 1, 2, 241.99, 483.98
        ]
    ]

    # Styling
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_data in rows:
        ws.append(row_data)

    for row in ws.iter_rows(min_row=2, max_row=len(rows)+1, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    # Auto adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    target_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "test_data.xlsx")
    wb.save(target_path)
    print(f"Successfully generated Excel test data at: {target_path}")

if __name__ == "__main__":
    generate_test_excel()
